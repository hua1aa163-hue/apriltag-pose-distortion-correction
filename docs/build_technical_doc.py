from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
SOURCE = DOCS / "反射投影回环畸变矫正-原理与实现.md"
OUTPUT = DOCS / "反射投影回环畸变矫正-技术文档-v1.1.docx"
RUN = ROOT / "artifacts" / "hardware-calibration-full" / "20260730-140908"
BUILD = DOCS / ".docx-build"
BUILD.mkdir(exist_ok=True)

FONT = "Microsoft YaHei"
BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
INK = RGBColor(0x0B, 0x25, 0x45)
MUTED = RGBColor(0x66, 0x6B, 0x73)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def set_run_font(run, size=None, bold=None, italic=None, color=None, font=FONT):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    if sum(widths) != TABLE_WIDTH_DXA:
        raise ValueError("table widths must sum to 9360 DXA")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    prefix = paragraph.add_run("第 ")
    set_run_font(prefix, size=9, color=MUTED)
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_begin, instr_text, fld_char_end])
    set_run_font(run, size=9, color=MUTED)
    suffix = paragraph.add_run(" 页")
    set_run_font(suffix, size=9, color=MUTED)


def add_numbering(document, kind):
    numbering = document.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    level.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    level.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    level.append(lvl_jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "260")
    tabs.append(tab)
    p_pr.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "540")
    indent.set(qn("w:hanging"), "270")
    p_pr.append(indent)
    level.append(p_pr)
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for attr in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{attr}"), FONT)
    r_pr.append(fonts)
    level.append(r_pr)
    abstract.append(level)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_num(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])
    p_pr.append(num_pr)


def configure_styles(document):
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    code = styles.add_style("CodeBlock", 1)
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    code.font.size = Pt(9)
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.right_indent = Inches(0.18)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(6)
    code.paragraph_format.line_spacing = 1.05

    caption = styles["Caption"]
    caption.font.name = FONT
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    caption.font.size = Pt(9)
    caption.font.color.rgb = MUTED
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_with_next = False


def add_inline_markup(paragraph, text, size=None, color=None):
    parts = re.split(r"(`[^`]+`|<https?://[^>]+>)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size=size or 10, color=color, font="Consolas")
            run.font.highlight_color = None
        elif part.startswith("<http") and part.endswith(">"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size=size or 10, color=BLUE)
            run.underline = True
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size, color=color)


def add_caption(document, text):
    p = document.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline_markup(p, text, size=9, color=MUTED)


def add_result_summary(document):
    table = document.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    set_cell_shading(table.cell(0, 0), CALLOUT)
    p = table.cell(0, 0).paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.2
    r = p.add_run("硬件实测通过  ")
    set_run_font(r, size=12, bold=True, color=INK)
    r = p.add_run("Gray Code 覆盖率 81.2%；189 点一次回环 RMS 0.583 camera px；最终自然逆映射约 1777×687。")
    set_run_font(r, size=10.5, color=INK)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def crop_overlay():
    source = RUN / "overlays" / "loop-00-overlay.png"
    target = BUILD / "loop-overlay-crop.png"
    with Image.open(source).convert("RGB") as image:
        pixels = image.load()
        xs, ys = [], []
        for y in range(0, image.height, 2):
            for x in range(0, image.width, 2):
                r, g, b = pixels[x, y]
                if g > 90 and g > r * 1.35 and g > b * 1.12:
                    xs.append(x)
                    ys.append(y)
        if not xs:
            raise RuntimeError("cannot locate green overlay")
        margin = 130
        box = (
            max(0, min(xs) - margin),
            max(0, min(ys) - margin),
            min(image.width, max(xs) + margin),
            min(image.height, max(ys) + margin),
        )
        image.crop(box).save(target)
    return target


def add_test_figures(document):
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(RUN / "patterns" / "final-corrected-27x7.png"), width=Inches(6.25))
    add_caption(document, "图 1  最终投影端预扭曲 27×7 点阵（1920×1080 画布）")

    crop = crop_overlay()
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(crop), width=Inches(6.15))
    add_caption(document, "图 2  相机回环识别叠加图（189 点全部识别，RMS 0.583 px）")

    table = document.add_table(rows=2, cols=2)
    set_table_geometry(table, [4680, 4680])
    labels = ("自动位置", "固定位置（120,260）")
    images = (
        RUN / "warp-sample" / "inverse-auto-position.png",
        RUN / "warp-sample" / "inverse-fixed-position.png",
    )
    for col, label in enumerate(labels):
        set_cell_shading(table.cell(0, col), LIGHT_BLUE)
        p = table.cell(0, col).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(label)
        set_run_font(r, size=10, bold=True, color=INK)
        p = table.cell(1, col).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(images[col]), width=Inches(3.02))
    add_caption(document, "图 3  任意完整 1920×1080 输入的两种放置方案；形状相同，仅整体位置不同")


def build_document():
    document = Document()
    configure_styles(document)
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    r = header.add_run("反射投影回环畸变矫正  |  技术报告与操作指南")
    set_run_font(r, size=8.5, color=MUTED)
    add_page_field(section.footer.paragraphs[0])

    # editorial_cover pattern, adapted for a Chinese technical manual.
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(64)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("反射投影系统")
    set_run_font(r, size=12, bold=True, color=BLUE)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("回环畸变矫正")
    set_run_font(r, size=30, bold=True, color=INK)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(30)
    r = p.add_run("原理、软件实现、GitHub 调研与硬件实测")
    set_run_font(r, size=14, color=DARK_BLUE)
    add_result_summary(document)
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(28)
    p.paragraph_format.space_after = Pt(4)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("MV-CS200-10GM  ·  1920×1080  ·  1777×687")
    set_run_font(r, size=10.5, bold=True, color=MUTED)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("版本 1.1  |  2026-07-30")
    set_run_font(r, size=10, color=MUTED)
    document.add_page_break()

    bullet_num = add_numbering(document, "bullet")
    decimal_num = add_numbering(document, "decimal")
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    start = next(i for i, line in enumerate(lines) if line.startswith("## 1."))
    lines = lines[start:]
    paragraph_buffer = []
    code_buffer = []
    in_code = False

    def flush_paragraph():
        nonlocal paragraph_buffer
        if not paragraph_buffer:
            return
        p = document.add_paragraph()
        text = " ".join(part.strip() for part in paragraph_buffer)
        add_inline_markup(p, text)
        paragraph_buffer = []

    def flush_code():
        nonlocal code_buffer
        if not code_buffer:
            return
        p = document.add_paragraph(style="CodeBlock")
        p.paragraph_format.keep_together = True
        set_cell = None
        r = p.add_run("\n".join(code_buffer))
        set_run_font(r, size=9, font="Consolas", color=INK)
        code_buffer = []

    for line in lines:
        if line.startswith("```"):
            if in_code:
                flush_code()
                in_code = False
            else:
                flush_paragraph()
                in_code = True
            continue
        if in_code:
            code_buffer.append(line)
            continue
        if not line.strip():
            flush_paragraph()
            continue
        if line.startswith("## "):
            flush_paragraph()
            heading = line[3:].strip()
            document.add_heading(heading, level=1)
            if heading.startswith("1. 结论"):
                add_result_summary(document)
            if heading.startswith("12. 实测产物"):
                add_test_figures(document)
            continue
        if line.startswith("### "):
            flush_paragraph()
            document.add_heading(line[4:].strip(), level=2)
            continue
        if line.startswith("- "):
            flush_paragraph()
            p = document.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.25
            apply_num(p, bullet_num)
            add_inline_markup(p, line[2:].strip())
            continue
        if re.match(r"^\d+\.\s", line):
            flush_paragraph()
            p = document.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.25
            apply_num(p, decimal_num)
            add_inline_markup(p, re.sub(r"^\d+\.\s+", "", line))
            continue
        paragraph_buffer.append(line)

    flush_paragraph()
    flush_code()

    props = document.core_properties
    props.title = "反射投影回环畸变矫正：原理、实现与实测"
    props.subject = "Gray Code 稠密映射、可配置点阵回环、海康 GigE 相机、Windows 桌面投图和文件夹批处理"
    props.author = "DistortionCorrection 项目组"
    props.keywords = "投影畸变矫正, Gray Code, 回环测试, 海康相机, 逆映射"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
